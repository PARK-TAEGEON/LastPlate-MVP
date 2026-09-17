import csv, io, json, math, zipfile, importlib.util
from datetime import date, datetime
from pydantic import BaseModel, ConfigDict, Field, ValidationError
from backend.bootstrap import NATIVE

MAX_BYTES=5*1024*1024
MAX_ROWS=5000
spec=importlib.util.spec_from_file_location('native_inventory_models',NATIVE/'inventory_risk/schemas/models.py')
native_models=importlib.util.module_from_spec(spec)
spec.loader.exec_module(native_models)


class UploadError(ValueError):
    def __init__(self,code,message,details=None):
        self.code,self.message,self.details=code,message,details or []
        super().__init__(message)


class MenuRow(BaseModel):
    model_config=ConfigDict(extra='forbid')
    date: date
    meal_type: str=Field(min_length=1)
    menu_name: str=Field(min_length=1,max_length=120)

class MonthlyMenuRow(BaseModel):
    model_config=ConfigDict(extra='forbid')
    date:date
    rice:str=Field(min_length=1,max_length=120)
    soup:str=Field(min_length=1,max_length=120)
    main:str=Field(min_length=1,max_length=120)
    side:str=Field(min_length=1,max_length=120)


def parse(content,filename,kind):
    if not content: raise UploadError('EMPTY_FILE','빈 파일입니다.')
    if len(content)>MAX_BYTES: raise UploadError('FILE_TOO_LARGE','파일은 5MB 이하로 올려주세요.')
    ext=filename.rsplit('.',1)[-1].lower()
    try:
        if ext=='csv':
            try: data=content.decode('utf-8-sig')
            except UnicodeDecodeError: data=content.decode('cp949')
            reader=csv.reader(io.StringIO(data),strict=True)
            matrix=[]
            for row in reader:
                matrix.append(row)
                if len(matrix)>MAX_ROWS+1 or len(row)>30:
                    raise UploadError('TOO_MANY_ROWS','최대 5,000행·30열까지 지원합니다.')
        elif ext=='xlsx':
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist())>30*1024*1024:
                    raise UploadError('FILE_TOO_LARGE','압축 해제 후 파일이 너무 큽니다.')
            book=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
            try:
                sheet=book.active
                if sheet.max_row>MAX_ROWS+1 or sheet.max_column>30:
                    raise UploadError('TOO_MANY_ROWS','최대 5,000행·30열까지 지원합니다.')
                matrix=[]
                for row in sheet.iter_rows():
                    if any(cell.data_type=='f' for cell in row):
                        raise UploadError('FORMULA_NOT_ALLOWED','수식 대신 계산된 값을 입력하세요.')
                    matrix.append([cell.value for cell in row])
            finally: book.close()
        else: raise UploadError('FILE_TYPE_ERROR','.csv 또는 .xlsx 파일만 지원합니다.')
    except UploadError:raise
    except Exception as exc:raise UploadError('MALFORMED_FILE','파일을 읽을 수 없습니다. 형식과 인코딩을 확인하세요.') from exc
    matrix=[r for r in matrix if any(v is not None and str(v).strip() for v in r)]
    if len(matrix)<2:raise UploadError('EMPTY_FILE','헤더 아래에 데이터 행이 필요합니다.')
    if len(matrix)>MAX_ROWS+1:raise UploadError('TOO_MANY_ROWS','최대 5,000행까지 지원합니다.')
    headers=[str(x or '').strip() for x in matrix[0]]
    if kind=='monthly-menu':
        aliases={'날짜':'date','밥':'rice','국':'soup','메인반찬':'main','사이드반찬':'side'}
        headers=[aliases.get(h,h) for h in headers]
    model=MonthlyMenuRow if kind=='monthly-menu' else MenuRow if kind=='menu' else native_models.InventoryLot
    required={k for k,v in model.model_fields.items() if v.is_required()}
    if len(headers)!=len(set(headers)) or not required.issubset(headers) or set(headers)-model.model_fields.keys():
        raise UploadError('INVALID_COLUMNS','필수 열 또는 열 이름을 확인하세요.',
                          [{'required':sorted(required),'allowed':list(model.model_fields),'received':headers}])
    records,seen=[],set()
    for index,cells in enumerate(matrix[1:],start=2):
        if len(cells)!=len(headers):raise UploadError('MALFORMED_ROW',f'{index}행의 열 수가 헤더와 다릅니다.')
        row={k:(v.strip() if isinstance(v,str) else v) for k,v in zip(headers,cells)}
        row={k:(v.date().isoformat() if isinstance(v,datetime) else v.isoformat() if isinstance(v,date) else v) for k,v in row.items()}
        row={k:v for k,v in row.items() if v not in ('',None)}
        if 'unit' in row and row['unit'] not in ('g','kg'):
            raise UploadError('UNIT_ERROR',f'{index}행: 지원하지 않는 단위입니다. g 또는 kg를 사용하세요.')
        try: normalized=model.model_validate(row).model_dump(mode='json')
        except ValidationError as exc:
            raise UploadError('ROW_VALIDATION_ERROR',f'{index}행의 값·날짜·필수 항목을 확인하세요.',
                              [{'field':'.'.join(map(str,e['loc'])),'message':e['msg']} for e in exc.errors()]) from exc
        identity=normalized['date'] if kind=='monthly-menu' else (normalized['date'],normalized['meal_type'],normalized['menu_name']) if kind=='menu' else json.dumps(normalized,sort_keys=True,ensure_ascii=False)
        if identity in seen:raise UploadError('DUPLICATE_ROW',f'{index}행이 앞선 데이터와 중복됩니다.')
        seen.add(identity);records.append(normalized)
    return records
